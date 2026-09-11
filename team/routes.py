# -*- coding: utf-8 -*-
"""Rotas do modulo de Gestao do Time (registradas sobre o app hospedeiro)."""
import fuso
import os
from functools import wraps
from datetime import datetime, date, timedelta

from flask import (render_template, request, redirect, url_for, flash, abort,
                   send_file, jsonify)
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename

from config import Config
from models import (db, User, Company, Competency, log_audit)
from team.models import (TeamMember, CompanyAssignment, Project, Milestone,
                         Activity, ClosingTemplateItem, Indicator, IndicatorResult,
                         AlertChannelSetting, AlertLog, ALERT_EVENTS,
                         KINDS, STATUSES, PRIORITIES, ensure_alert_defaults)
from team import engine
from team import alerts as team_alerts

EVIDENCE_DIR = os.path.join(Config.UPLOAD_DIR, "_evidencias")
ALLOWED_EVIDENCE = {".pdf", ".xlsx", ".xlsm", ".png", ".jpg", ".jpeg",
                    ".docx", ".pptx", ".csv", ".txt", ".msg"}


def register_team_routes(app):

    # ------------------------------------------------------------------
    # Guard de perfil: modulo do time (controladoria, admin, profissional).
    # O perfil 'profissional' entra, mas so enxerga o proprio painel (escopo).
    # ------------------------------------------------------------------
    def team_required(f):
        @wraps(f)
        @login_required
        def wrap(*a, **k):
            if not current_user.is_team:
                abort(403)
            return f(*a, **k)
        return wrap

    def controladoria_required(f):
        """Gestão da área (cronograma, etc.): admin/controladoria só."""
        @wraps(f)
        @login_required
        def wrap(*a, **k):
            if not current_user.is_controladoria:
                abort(403)
            return f(*a, **k)
        return wrap

    def _current_competency():
        return app._current_competency()

    def _current_member():
        return TeamMember.query.filter_by(user_id=current_user.id).first()

    def _scoped_member_id():
        """Se o usuario e 'profissional', devolve o id do proprio membro para
        filtrar as visoes. -1 quando ele nao tem membro vinculado (ve nada).
        None para admin/controladoria (ve tudo)."""
        if not current_user.is_profissional:
            return None
        m = _current_member()
        return m.id if m else -1

    def _members(active_only=True):
        """Lista de membros, ja escopada quando o usuario e profissional."""
        q = TeamMember.query
        if active_only:
            q = q.filter_by(active=True)
        sid = _scoped_member_id()
        if sid is not None:
            q = q.filter(TeamMember.id == sid)
        return q.order_by(TeamMember.sort_order, TeamMember.name).all()

    def _member_map():
        return {m.id: m for m in TeamMember.query.all()}

    # -------- liderança de projeto + delegação por férias --------------------
    def _absence_atual(member, ref=None):
        """Ausência aprovada da pessoa cobrindo a data (ou None)."""
        from team.models_workflow import Absence
        if not member:
            return None
        ref = ref or fuso.hoje()
        return (Absence.query.filter_by(member_id=member.id, status="aprovada")
                .filter(Absence.start_date <= ref, Absence.end_date >= ref).first())

    def _membro_em_ferias(member, ref=None):
        return _absence_atual(member, ref) is not None

    def _team_manager(exclude_id=None):
        q = TeamMember.query.filter_by(is_manager=True, active=True)
        if exclude_id:
            q = q.filter(TeamMember.id != exclude_id)
        return q.order_by(TeamMember.sort_order).first()

    def _resolve_delegacao(member, campo, ref=None):
        """Segue a cadeia de delegação (aprova/andamento) até quem NÃO está de férias.

        `campo` = 'aprova_delegado' ou 'andamento_delegado'. Se a pessoa está fora e
        não nomeou ninguém, escala para o gestor do time. Evita ciclos."""
        seen, hops, cur = set(), 0, member
        while cur and cur.id not in seen and hops < 6:
            seen.add(cur.id)
            ab = _absence_atual(cur, ref)
            if not ab:
                return cur                    # disponível: é o responsável efetivo
            cur = getattr(ab, campo) or _team_manager(exclude_id=cur.id)
            hops += 1
        return cur

    def _projeto_gestor(p):
        """Gestor do projeto: o definido no projeto, senão o gestor do time."""
        if p and p.manager_member_id:
            g = db.session.get(TeamMember, p.manager_member_id)
            if g:
                return g
        return _team_manager()

    def _projeto_lider(p, ref=None):
        """Quem TOCA o projeto agora (andamento) — considera férias/delegação."""
        if not p or not p.owner:
            return _projeto_gestor(p)
        return _resolve_delegacao(p.owner, "andamento_delegado", ref)

    def _lider_em_ferias(p, ref=None):
        return bool(p and p.owner and _membro_em_ferias(p.owner, ref))

    def _aprovadores_projeto(p, ref=None):
        """user_ids que podem aprovar realinhamentos deste projeto agora
        (líder e gestor, seguindo a delegação de aprovação quando de férias)."""
        ids = set()
        for base in (p.owner if p else None, _projeto_gestor(p)):
            if base:
                efetivo = _resolve_delegacao(base, "aprova_delegado", ref)
                if efetivo and efetivo.user_id:
                    ids.add(efetivo.user_id)
        return ids

    def _pode_gerir_projeto(p):
        """Admin, ou o líder/gestor efetivo (com delegação). Sem líder/gestor
        definido, cai para controladoria (evita travar a gestão)."""
        if current_user.is_admin:
            return True
        aprovs = _aprovadores_projeto(p)
        if aprovs:
            return current_user.id in aprovs
        return current_user.is_controladoria

    def _entity_project(rev_type, entity_id):
        if rev_type == "project":
            return db.session.get(Project, entity_id)
        if rev_type == "milestone":
            m = db.session.get(Milestone, entity_id)
            return m.project if m else None
        if rev_type == "activity":
            a = db.session.get(Activity, entity_id)
            return a.project if a else None
        return None

    def _pode_aprovar_revisao(rev_type, entity_id):
        proj = _entity_project(rev_type, entity_id)
        if proj:
            return _pode_gerir_projeto(proj)
        return current_user.is_controladoria    # atividade sem projeto: gestor de área

    def _aplicar_revisao(rev):
        if rev.entity_type == "project":
            p = db.session.get(Project, rev.entity_id)
            if p:
                p.target_date = rev.new_date
        elif rev.entity_type == "milestone":
            m = db.session.get(Milestone, rev.entity_id)
            if m:
                m.due_date = rev.new_date
        elif rev.entity_type == "activity":
            a = db.session.get(Activity, rev.entity_id)
            if a:
                a.due_date = rev.new_date
                a.due_provisional = False

    def _revisar_prazo(entity_type, entity_id, old_date, label, back_url):
        """Aplica direto (líder/gestor) ou cria solicitação de aprovação."""
        from team.models_workflow import DeadlineRevision
        from models import notify
        nova = _date(request.form.get("new_date"))
        motivo = (request.form.get("reason") or "").strip()
        if not nova or not motivo:
            flash("Informe a nova data e a justificativa.", "danger")
            return redirect(back_url)
        if _pode_aprovar_revisao(entity_type, entity_id):
            rev = DeadlineRevision(
                entity_type=entity_type, entity_id=entity_id, old_date=old_date,
                new_date=nova, reason=motivo, created_by=current_user.id,
                status="aplicada", approved_by=current_user.id,
                approved_at=datetime.utcnow())
            db.session.add(rev)
            _aplicar_revisao(rev)
            db.session.commit()
            flash("Prazo realinhado e registrado no histórico.", "success")
        else:
            rev = DeadlineRevision(
                entity_type=entity_type, entity_id=entity_id, old_date=old_date,
                new_date=nova, reason=motivo, created_by=current_user.id,
                status="solicitada", requested_by=current_user.id)
            db.session.add(rev)
            db.session.commit()
            proj = _entity_project(entity_type, entity_id)
            # notifica quem PODE aprovar agora (líder/gestor efetivos c/ delegação)
            alvos = set(_aprovadores_projeto(proj)) if proj else set()
            if not alvos:      # sem projeto ou sem líder/gestor: cai nos gestores
                alvos = {u.id for u in User.query.filter(
                    User.role.in_(["controladoria", "admin"])).all()}
            de = old_date.strftime("%d/%m") if old_date else "—"
            for uid in alvos:
                notify(uid, "Realinhamento de prazo a aprovar",
                       f"{label}: {de} → {nova.strftime('%d/%m/%Y')} — {motivo[:80]}",
                       kind="prazo", url=back_url)
            flash("Realinhamento enviado para aprovação do líder/gestor.", "warning")
        return redirect(back_url)

    # ==================================================================
    # PAINEL DO DIA
    # ==================================================================
    @app.route("/time")
    @team_required
    def team_hoje():
        ref = fuso.hoje()
        members = _members()
        sel = request.args.get("member_id", type=int)
        tipo = request.args.get("tipo")     # 'projeto' | 'fechamento' | None
        if tipo not in ("projeto", "fechamento"):
            tipo = None
        panel = engine.day_panel(ref=ref, member_id=sel, kind_filter=tipo)
        comp = _current_competency()
        fr = engine.farol(comp, ref) if comp else None
        mfarol = engine.member_farol(ref=ref)
        mm = _member_map()
        # resumo por membro (carga em aberto)
        board = []
        for m in members:
            mf = mfarol.get(m.id, {"total": 0, "atrasada": 0,
                                   "vence_hoje": 0, "aberta": 0})
            board.append({"member": m, "f": mf})
        rg = engine.regua(comp, ref, member_id=sel) if comp else None
        # barra de cada pessoa proporcional à maior carga do time
        maior = max([b["f"]["aberta"] for b in board] + [1])
        return render_template("team/hoje.html", panel=panel, members=members,
                               sel=sel, tipo=tipo, comp=comp, farol=fr, board=board,
                               mm=mm, today=ref, regua=rg, maior_carga=maior,
                               du_entre=engine.business_days_between)

    # ==================================================================
    # ATIVIDADES (motor unico) — lista, kanban, CRUD
    # ==================================================================
    @app.route("/atividades")
    @team_required
    def team_atividades():
        view = request.args.get("view", "tabela")
        group = request.args.get("group")     # 'status' agrupa a tabela por status
        f_kind = request.args.get("kind")
        f_member = request.args.get("member_id", type=int)
        f_status = request.args.get("status")
        f_comp = request.args.get("competency_id", type=int)
        f_project = request.args.get("project_id", type=int)
        q = Activity.query
        sid = _scoped_member_id()
        if sid is not None:               # profissional: so as proprias
            q = q.filter(Activity.member_id == sid)
        if f_kind:
            q = q.filter(Activity.kind == f_kind)
        if f_member and sid is None:
            q = q.filter(Activity.member_id == f_member)
        if f_comp:
            q = q.filter(Activity.competency_id == f_comp)
        if f_project:
            q = q.filter(Activity.project_id == f_project)
        acts = q.order_by(Activity.due_date.is_(None), Activity.due_date,
                          Activity.sort_order, Activity.id).all()
        ref = fuso.hoje()
        if f_status:
            acts = [a for a in acts if a.effective_status(ref) == f_status]
        # colunas do kanban por status efetivo
        columns = {"atrasada": [], "vence_hoje": [], "pendente": [], "em_andamento": [],
                   "aguardando": [], "bloqueada": [], "concluida": [], "cancelada": []}
        for a in acts:
            columns.setdefault(a.effective_status(ref), []).append(a)
        # agrupamento por STATUS efetivo (colapsável), quando pedido
        st_ordem = ["atrasada", "vence_hoje", "pendente", "em_andamento",
                    "aguardando", "bloqueada", "concluida", "cancelada"]
        st_labels = {"atrasada": "Atrasada", "vence_hoje": "Vence hoje",
                     "pendente": "Pendente", "em_andamento": "Em andamento",
                     "aguardando": "Aguardando insumo", "bloqueada": "Bloqueada",
                     "concluida": "Concluída", "cancelada": "Cancelada"}
        grupos_status = []
        if group == "status":
            for k in st_ordem:
                rows = [a for a in acts if a.effective_status(ref) == k]
                if rows:
                    grupos_status.append((k, st_labels.get(k, k), rows))
        return render_template("team/atividades.html", acts=acts, view=view,
                               columns=columns, members=_members(), group=group,
                               grupos_status=grupos_status,
                               competencies=Competency.query.order_by(
                                   Competency.year.desc(), Competency.month.desc()).all(),
                               projects=Project.query.order_by(Project.name).all(),
                               kinds=KINDS, statuses=STATUSES, mm=_member_map(),
                               f_kind=f_kind, f_member=f_member, f_status=f_status,
                               f_comp=f_comp, f_project=f_project, today=ref)

    @app.route("/atividade/nova", methods=["GET", "POST"])
    @team_required
    def team_atividade_new():
        if request.method == "POST":
            a = _apply_activity_form(Activity(status="pendente", origin="manual",
                                              created_by=current_user.id), request.form)
            db.session.add(a)
            db.session.commit()
            log_audit(current_user.id, "atividade_criada", "activity", str(a.id))
            flash("Atividade criada.", "success")
            return redirect(url_for("team_atividade", aid=a.id))
        return render_template("team/atividade_form.html", a=None,
                               members=_members(), companies=_companies(),
                               competencies=_comps(), projects=_projects(),
                               kinds=KINDS, priorities=PRIORITIES,
                               preset_project=request.args.get("project_id", type=int))

    @app.route("/atividade/<int:aid>")
    @team_required
    def team_atividade(aid):
        a = db.session.get(Activity, aid) or abort(404)
        from team.models_workflow import DeadlineRevision
        todas = (DeadlineRevision.query
                 .filter_by(entity_type="activity", entity_id=a.id)
                 .order_by(DeadlineRevision.created_at.desc()).all())
        revisoes = [r for r in todas if r.status == "aplicada"]
        pendentes = [r for r in todas if r.status == "solicitada"]
        return render_template("team/atividade_detail.html", a=a, mm=_member_map(),
                               revisoes=revisoes, pendentes=pendentes,
                               pode_aprovar=_pode_aprovar_revisao("activity", a.id),
                               users={u.id: u for u in User.query.all()},
                               today=fuso.hoje())

    @app.route("/atividade/<int:aid>/editar", methods=["GET", "POST"])
    @team_required
    def team_atividade_edit(aid):
        a = db.session.get(Activity, aid) or abort(404)
        if request.method == "POST":
            _apply_activity_form(a, request.form)
            db.session.commit()
            log_audit(current_user.id, "atividade_editada", "activity", str(a.id))
            flash("Atividade atualizada.", "success")
            return redirect(url_for("team_atividade", aid=a.id))
        return render_template("team/atividade_form.html", a=a, members=_members(),
                               companies=_companies(), competencies=_comps(),
                               projects=_projects(), kinds=KINDS, priorities=PRIORITIES)

    @app.route("/atividade/<int:aid>/status", methods=["POST"])
    @team_required
    def team_atividade_status(aid):
        a = db.session.get(Activity, aid) or abort(404)
        new = request.form.get("status")
        if new in STATUSES:
            a.status = new
            if new == "concluida" and not a.done_at:
                a.done_at = datetime.utcnow()
                a.done_by = current_user.id
            if new != "concluida":
                a.done_at = None
                a.done_by = None
            db.session.commit()
            log_audit(current_user.id, f"atividade_{new}", "activity", str(a.id))
        if request.form.get("ajax"):
            return jsonify(ok=True, status=a.effective_status())
        flash("Status atualizado.", "success")
        return redirect(request.referrer or url_for("team_atividade", aid=a.id))

    @app.route("/atividade/<int:aid>/excluir", methods=["POST"])
    @team_required
    def team_atividade_delete(aid):
        a = db.session.get(Activity, aid) or abort(404)
        db.session.delete(a)
        db.session.commit()
        log_audit(current_user.id, "atividade_excluida", "activity", str(aid))
        flash("Atividade excluída.", "success")
        return redirect(url_for("team_atividades"))

    @app.route("/atividade/<int:aid>/reatribuir", methods=["POST"])
    @team_required
    def team_atividade_reassign(aid):
        """Troca só o responsável — sem tocar em prazo, empresa ou status."""
        a = db.session.get(Activity, aid) or abort(404)
        a.member_id = _int(request.form.get("member_id"))
        db.session.commit()
        log_audit(current_user.id, "atividade_reatribuida", "activity", str(aid))
        flash("Atividade reatribuída.", "success")
        return redirect(request.referrer or url_for("team_atividades"))

    def _apply_activity_form(a, form):
        a.title = (form.get("title") or "").strip() or a.title or "Sem título"
        a.description = form.get("description") or None
        a.kind = form.get("kind") if form.get("kind") in KINDS else (a.kind or "spot")
        a.priority = form.get("priority") if form.get("priority") in PRIORITIES else "media"
        a.member_id = _int(form.get("member_id"))
        a.company_id = _int(form.get("company_id"))
        a.project_id = _int(form.get("project_id"))
        a.competency_id = _int(form.get("competency_id"))
        # atividade de projeto: sempre tipo 'projeto' e sem competência (não é fechamento)
        if a.project_id:
            a.kind = "projeto"
            a.competency_id = None
        a.start_date = _date(form.get("start_date"))
        a.due_date = _date(form.get("due_date"))
        st = form.get("status")
        if st in STATUSES:
            a.status = st
            if st == "concluida" and not a.done_at:
                a.done_at = datetime.utcnow()
                a.done_by = current_user.id
        return a

    # ==================================================================
    # AGENDA DE FECHAMENTO (template + geracao)
    # ==================================================================
    @app.route("/agenda")
    @team_required
    def team_agenda():
        items = (ClosingTemplateItem.query
                 .order_by(ClosingTemplateItem.sort_order, ClosingTemplateItem.id).all())
        comp = _current_competency()
        comps = _comps()
        # preview do que cada item viraria na competencia atual
        preview = []
        if comp:
            for it in items:
                rule = {"base": it.due_base, "offset": it.due_offset,
                        "codes": it.insumo_codes}
                due, prov = engine.compute_due(rule, comp)
                preview.append({"item": it, "due": due, "prov": prov})
        gen_count = Activity.query.filter_by(
            competency_id=comp.id, origin="template").count() if comp else 0
        return render_template("team/agenda.html", items=items, comp=comp,
                               comps=comps, preview=preview, members=_members(),
                               companies=_companies(), gen_count=gen_count)

    @app.route("/agenda/item", methods=["POST"])
    @team_required
    def team_agenda_item():
        it = ClosingTemplateItem(
            title=(request.form.get("title") or "").strip() or "Novo item",
            kind=request.form.get("kind") if request.form.get("kind") in KINDS else "fechamento",
            member_id=_int(request.form.get("member_id")),
            company_id=_int(request.form.get("company_id")),
            deliverable=request.form.get("deliverable") or None,
            priority=request.form.get("priority") if request.form.get("priority") in PRIORITIES else "media",
            due_base=request.form.get("due_base") or "deadline",
            due_offset=_int(request.form.get("due_offset")) or 0,
            auto_metric=bool(request.form.get("auto_metric")),
            sort_order=_int(request.form.get("sort_order")) or 100)
        codes = request.form.get("insumo_codes") or ""
        it.insumo_codes = [c.strip() for c in codes.split(",") if c.strip()]
        db.session.add(it)
        db.session.commit()
        log_audit(current_user.id, "template_item_criado", "closing_template", str(it.id))
        flash("Item da agenda criado.", "success")
        return redirect(url_for("team_agenda"))

    @app.route("/agenda/item/<int:iid>/toggle", methods=["POST"])
    @team_required
    def team_agenda_toggle(iid):
        it = db.session.get(ClosingTemplateItem, iid) or abort(404)
        it.active = not it.active
        db.session.commit()
        return redirect(url_for("team_agenda"))

    @app.route("/agenda/item/<int:iid>/excluir", methods=["POST"])
    @team_required
    def team_agenda_item_delete(iid):
        it = db.session.get(ClosingTemplateItem, iid) or abort(404)
        db.session.delete(it)
        db.session.commit()
        flash("Item removido.", "success")
        return redirect(url_for("team_agenda"))

    @app.route("/agenda/gerar", methods=["POST"])
    @team_required
    def team_agenda_generate():
        cid = _int(request.form.get("competency_id"))
        comp = db.session.get(Competency, cid) if cid else _current_competency()
        if not comp:
            flash("Selecione uma competência.", "warning")
            return redirect(url_for("team_agenda"))
        criadas, atualizadas, removidas, preservadas = \
            engine.generate_closing_activities(comp, created_by=current_user.id)
        engine.resolve_auto_metrics(comp)
        log_audit(current_user.id, "agenda_gerada", "competency",
                  f"{comp.key} +{criadas}/~{atualizadas}/-{removidas}")
        flash(f"Agenda de {comp.label}: {criadas} criada(s), {atualizadas} atualizada(s), "
              f"{removidas} removida(s), {preservadas} concluída(s) preservada(s).",
              "success")
        return redirect(url_for("team_atividades", competency_id=comp.id, kind="fechamento"))

    @app.route("/agenda/sincronizar", methods=["POST"])
    @team_required
    def team_agenda_sync():
        comp = _current_competency()
        if not comp:
            flash("Nenhuma competência aberta.", "warning")
            return redirect(url_for("team_agenda"))
        upd = engine.refresh_provisional_due_dates(comp)
        done = engine.resolve_auto_metrics(comp)
        flash(f"Sincronizado: {upd} prazo(s) recalculado(s), {done} atividade(s) "
              f"concluída(s) por evidência de fechamento.", "success")
        return redirect(url_for("team_agenda"))

    # ==================================================================
    # CRONOGRAMA DE FECHAMENTO (cadastro de itens + geração mês/período/ano)
    # ==================================================================
    def _cron_item_from_form(it, form, p="", novo=False):
        """Aplica os campos do formulário ao item. `p` = prefixo do campo (lote:
        't12_' para o item 12, 'n3_' para a 3ª linha nova/duplicada)."""
        def g(k, d=None):
            return form.get(p + k, d)
        it.title = (g("title") or "").strip() or it.title or "Nova atividade"
        if (p + "kind") in form:
            it.kind = g("kind") if g("kind") in KINDS else "fechamento"
        elif not it.kind:
            it.kind = "fechamento"
        # escopo: 'geral' | 'todas' | 'empresas' (+ lista 'empresas') | <id> (formato antigo)
        escopo = (g("escopo") or "geral").strip()
        empresas = [e for e in form.getlist(p + "empresas") if e]
        if escopo == "todas":
            it.per_company = True
            it.company_ids = []
        elif escopo == "empresas" or (escopo not in ("geral", "todas") and escopo.isdigit()):
            if escopo.isdigit():
                empresas = [escopo] + empresas
            it.per_company = False
            it.company_ids = empresas             # sem nenhuma marcada = vira geral
        else:
            it.per_company = False
            it.company_ids = []
        # responsável (usado no geral e como fallback nas empresas)
        it.member_id = _int(g("member_id"))
        if (p + "deliverable") in form:
            it.deliverable = g("deliverable") or None
        it.priority = g("priority") if g("priority") in PRIORITIES else (it.priority or "media")
        if (p + "active") in form or (p + "active_presente") in form:
            it.active = g("active") in ("1", "on", "true")
        # prazo: o cronograma trabalha em Nº dia útil do mês seguinte. Itens antigos
        # com outra base (ex.: 5º DU −1) só mudam se o número for de fato editado.
        off = _int(g("due_offset"))
        if off is not None and (novo or it.due_base == "fixed_bd" or off != it.due_offset):
            it.due_base = "fixed_bd"
            it.due_offset = max(off, 1)
        elif novo and it.due_offset is None:
            it.due_base, it.due_offset = "fixed_bd", 5
        if _int(g("sort_order")):
            it.sort_order = _int(g("sort_order"))
        elif not it.sort_order:
            it.sort_order = 100

    def _exclui_item_cronograma(it):
        """Remove o item e DESVINCULA as atividades geradas por ele. O SQLite
        reaproveita o id apagado: sem isso, um item criado depois herdaria as
        atividades (e o histórico) do excluído. Desvinculadas, as concluídas
        ficam como histórico e as em aberto saem na próxima geração."""
        Activity.query.filter_by(template_id=it.id).update(
            {Activity.template_id: None}, synchronize_session=False)
        db.session.delete(it)

    @app.route("/cronograma/salvar", methods=["POST"])
    @controladoria_required
    def team_cronograma_salvar():
        """Salvar tudo: grava de uma vez as linhas alteradas, cria as duplicadas
        e remove as marcadas para exclusão."""
        import re as _re
        f = request.form
        ids = sorted({int(m.group(1)) for k in f for m in [_re.match(r"t(\d+)_", k)] if m})
        novos = sorted({int(m.group(1)) for k in f for m in [_re.match(r"n(\d+)_", k)] if m})
        alterados = excluidos = criados = 0
        for iid in ids:
            it = db.session.get(ClosingTemplateItem, iid)
            if not it:
                continue
            if f.get(f"t{iid}_excluir") == "1":
                _exclui_item_cronograma(it)
                excluidos += 1
            elif f.get(f"t{iid}_mudou") == "1":
                _cron_item_from_form(it, f, f"t{iid}_")
                alterados += 1
        db.session.flush()
        for k in novos:
            p = f"n{k}_"
            if f.get(p + "excluir") == "1" or not (f.get(p + "title") or "").strip():
                continue
            origem = db.session.get(ClosingTemplateItem, _int(f.get(p + "origem")) or 0)
            it = ClosingTemplateItem(title="Nova atividade")
            if origem:                                   # duplicata: herda o que não está na tela
                it.kind, it.deliverable = origem.kind, origem.deliverable
                it.due_base, it.due_offset = origem.due_base, origem.due_offset
                it.insumo_codes_json, it.auto_metric = origem.insumo_codes_json, origem.auto_metric
                it.sort_order = origem.sort_order
            _cron_item_from_form(it, f, p, novo=not origem)
            db.session.add(it)
            criados += 1
        db.session.commit()
        log_audit(current_user.id, "cronograma_salvo", "closing_template",
                  f"~{alterados} +{criados} -{excluidos}")
        partes = []
        if alterados:
            partes.append(f"{alterados} alterada(s)")
        if criados:
            partes.append(f"{criados} criada(s)")
        if excluidos:
            partes.append(f"{excluidos} removida(s)")
        flash("Cronograma salvo: " + ", ".join(partes) + "." if partes else "Nada a salvar.",
              "success" if partes else "info")
        return redirect(url_for("team_cronograma"))

    @app.route("/cronograma")
    @controladoria_required
    def team_cronograma():
        items = (ClosingTemplateItem.query
                 .order_by(ClosingTemplateItem.sort_order, ClosingTemplateItem.id).all())
        hoje = fuso.hoje()
        return render_template("team/cronograma.html", items=items,
                               members=_members(), companies=_companies(),
                               kinds=KINDS, priorities=PRIORITIES,
                               ano_atual=hoje.year, mes_atual=hoje.month,
                               n_empresas=CompanyAssignment.query.count())

    @app.route("/cronograma/item", methods=["POST"])
    @controladoria_required
    def team_cronograma_item_new():
        it = ClosingTemplateItem(title="Nova atividade")
        _cron_item_from_form(it, request.form, novo=True)
        db.session.add(it)
        db.session.commit()
        log_audit(current_user.id, "cronograma_item_criado", "closing_template", str(it.id))
        flash("Atividade adicionada ao cronograma.", "success")
        return redirect(url_for("team_cronograma"))

    @app.route("/cronograma/item/<int:iid>", methods=["POST"])
    @controladoria_required
    def team_cronograma_item_edit(iid):
        it = db.session.get(ClosingTemplateItem, iid) or abort(404)
        _cron_item_from_form(it, request.form)
        db.session.commit()
        flash("Atividade do cronograma atualizada.", "success")
        return redirect(url_for("team_cronograma"))

    @app.route("/cronograma/item/<int:iid>/toggle", methods=["POST"])
    @controladoria_required
    def team_cronograma_item_toggle(iid):
        it = db.session.get(ClosingTemplateItem, iid) or abort(404)
        it.active = not it.active
        db.session.commit()
        return redirect(url_for("team_cronograma"))

    @app.route("/cronograma/item/<int:iid>/excluir", methods=["POST"])
    @controladoria_required
    def team_cronograma_item_delete(iid):
        it = db.session.get(ClosingTemplateItem, iid) or abort(404)
        _exclui_item_cronograma(it)
        db.session.commit()
        flash("Atividade removida do cronograma.", "success")
        return redirect(url_for("team_cronograma"))

    @app.route("/cronograma/gerar", methods=["POST"])
    @controladoria_required
    def team_cronograma_gerar():
        modo = request.form.get("modo") or "mes"
        ano = _int(request.form.get("ano")) or fuso.hoje().year
        pares = []
        if modo == "ano":
            pares = [(ano, m) for m in range(1, 13)]
        elif modo == "periodo":
            ini = _int(request.form.get("mes_ini")) or 1
            fim = _int(request.form.get("mes_fim")) or 12
            if fim < ini:
                ini, fim = fim, ini
            pares = [(ano, m) for m in range(ini, fim + 1)]
        else:  # mes
            m = _int(request.form.get("mes")) or fuso.hoje().month
            pares = [(ano, m)]
        if not ClosingTemplateItem.query.filter_by(active=True).count():
            flash("Cadastre ao menos uma atividade ativa antes de gerar.", "warning")
            return redirect(url_for("team_cronograma"))
        criadas, atualizadas, removidas, preservadas, ncomp = \
            engine.generate_for_range(pares, created_by=current_user.id)
        log_audit(current_user.id, "cronograma_gerado", "competency",
                  f"{modo} {ano}: +{criadas}/~{atualizadas}/-{removidas}")
        flash(f"Cronograma sincronizado em {ncomp} competência(s): "
              f"{criadas} criada(s), {atualizadas} atualizada(s), "
              f"{removidas} removida(s), {preservadas} concluída(s) preservada(s).",
              "success")
        return redirect(url_for("team_atividades", kind="fechamento"))

    # ==================================================================
    # INDICADORES / METAS
    # ==================================================================
    @app.route("/indicadores")
    @team_required
    def team_indicadores():
        from team.models_workflow import Panel
        f_panel = request.args.get("panel_id", type=int)
        q = Indicator.query.filter_by(active=True)
        # profissional so ve os paineis dele (o proprio + o que segue)
        meus_paineis = None
        sid = _scoped_member_id()
        if sid is not None:
            me = _current_member()
            ids = set()
            if me:
                if me.panel_id:
                    ids.add(me.panel_id)
                proprio = Panel.query.filter_by(kind="pessoal",
                                                owner_member_id=me.id).first()
                if proprio:
                    ids.add(proprio.id)
            meus_paineis = ids or {-1}
            q = q.filter(Indicator.panel_id.in_(meus_paineis))
        if f_panel and (meus_paineis is None or f_panel in meus_paineis):
            q = q.filter(Indicator.panel_id == f_panel)
        inds = q.order_by(Indicator.panel_id, Indicator.seq).all()
        # indicadores de membros inativos somem daqui (ficam só no cadastro)
        inds = [i for i in inds if not (i.member and not i.member.active)]
        # agrupa por painel
        groups = {}
        for i in inds:
            groups.setdefault(i.panel_id, []).append(i)
        pq = Panel.query.filter_by(active=True)
        if meus_paineis is not None:
            pq = pq.filter(Panel.id.in_(meus_paineis))
        panels = pq.order_by(Panel.sort_order, Panel.name).all()
        pmap = {p.id: p for p in Panel.query.all()}
        from team.models import IndicatorDef
        defs = (IndicatorDef.query.filter_by(active=True)
                .order_by(IndicatorDef.sort_order, IndicatorDef.title).all())
        totais = {pid: round(sum(i.weight or 0 for i in lst), 2)
                  for pid, lst in groups.items()}
        return render_template("team/indicadores.html", groups=groups, totais=totais,
                               panels=panels, pmap=pmap, f_panel=f_panel,
                               defs=defs, escopo=(sid is not None))

    def _peso(v):
        """'25', '25%', '12,5' -> 25.0 / 12.5; vazio -> None; fora de 0..100 -> erro."""
        v = (v or "").strip().replace("%", "").replace(",", ".")
        if not v:
            return None
        p = float(v)
        if not 0 <= p <= 100:
            raise ValueError("peso fora de 0 a 100")
        return p

    def _peso_ou_nada(v):
        try:
            return _peso(v)
        except ValueError:
            flash("Peso ignorado: use um número de 0 a 100.", "warning")
            return None

    @app.route("/indicador/<int:iid>/peso", methods=["POST"])
    @controladoria_required
    def team_indicador_peso(iid):
        i = db.session.get(Indicator, iid) or abort(404)
        try:
            i.weight = _peso(request.form.get("weight"))
        except ValueError:
            return jsonify(ok=False, erro="Use um número de 0 a 100."), 400
        db.session.commit()
        log_audit(current_user.id, "indicador_peso", "indicator", f"{i.id}: {i.weight}")
        irmaos = Indicator.query.filter_by(panel_id=i.panel_id, active=True).all()
        total = sum(x.weight or 0 for x in irmaos)
        return jsonify(ok=True, weight=i.weight, total=round(total, 2))

    @app.route("/indicador/nova", methods=["POST"])
    @team_required
    def team_indicador_new():
        from team.models_workflow import Panel
        from team.models import IndicatorDef
        pid = _int(request.form.get("panel_id"))
        pan = db.session.get(Panel, pid) if pid else None
        # meta de painel pessoal -> vinculada ao dono; equipe -> sem dono unico
        member_id = pan.owner_member_id if (pan and pan.kind == "pessoal") else None

        # o indicador vem do catálogo (definição). Se escolheu "novo", cria a def.
        did = _int(request.form.get("indicator_def_id"))
        d = db.session.get(IndicatorDef, did) if did else None
        if not d:
            titulo = (request.form.get("new_title") or
                      request.form.get("title") or "").strip()
            if not titulo:
                flash("Escolha um indicador do catálogo ou informe um novo.", "danger")
                return redirect(url_for("team_indicadores", panel_id=pid))
            # reaproveita definição igual, se já existir
            d = IndicatorDef.query.filter(
                db.func.lower(IndicatorDef.title) == titulo.lower()).first()
            if not d:
                d = IndicatorDef(
                    title=titulo,
                    dimension=request.form.get("dimension") or "PRAZO",
                    target_type=request.form.get("target_type") or "manual",
                    rational=request.form.get("rational") or None)
                db.session.add(d)
                db.session.flush()

        # próximo seq dentro do painel
        prox = (db.session.query(db.func.max(Indicator.seq))
                .filter(Indicator.panel_id == pid).scalar() or 0) + 1
        i = Indicator(
            indicator_def_id=d.id,
            panel_id=pid,
            member_id=member_id,
            seq=_int(request.form.get("seq")) or prox,
            title=d.title,                       # denormalizado p/ exibição
            dimension=d.dimension,
            target_label=request.form.get("target_label") or None,   # a META do painel
            weight=_peso_ou_nada(request.form.get("weight")),
            target_type=d.target_type,
            rational=request.form.get("rational") or d.rational,
            auto_source=d.auto_source,
            year=_int(request.form.get("year")) or 2026)
        db.session.add(i)
        db.session.commit()
        flash("Indicador adicionado ao painel.", "success")
        if request.form.get("next"):
            return redirect(request.form.get("next"))
        return redirect(url_for("team_indicador", iid=i.id))

    # ---------------- Catálogo de indicadores (definições reutilizáveis) -------
    @app.route("/indicadores/catalogo")
    @controladoria_required
    def team_indicador_catalogo():
        from team.models import IndicatorDef
        defs = (IndicatorDef.query
                .order_by(IndicatorDef.active.desc(), IndicatorDef.sort_order,
                          IndicatorDef.title).all())
        return render_template("team/catalogo.html", defs=defs,
                               dimensions=["PRAZO", "RESULTADO", "PROCESSO", "QUALIDADE"],
                               target_types=[("manual", "Manual"),
                                             ("prazo_du", "Prazo (dia útil)"),
                                             ("contagem", "Contagem"),
                                             ("data_marco", "Data / marco"),
                                             ("projeto", "Projeto")])

    def _def_from_form(d, form):
        d.title = (form.get("title") or "").strip() or d.title or "Novo indicador"
        d.dimension = form.get("dimension") or "PRAZO"
        d.target_type = form.get("target_type") or "manual"
        d.unit = form.get("unit") or None
        d.rational = form.get("rational") or None
        d.sort_order = _int(form.get("sort_order")) or d.sort_order or 100

    @app.route("/catalogo/indicador", methods=["POST"])
    @controladoria_required
    def team_indicador_def_new():
        from team.models import IndicatorDef
        d = IndicatorDef(title="Novo indicador")
        _def_from_form(d, request.form)
        db.session.add(d)
        db.session.commit()
        flash("Indicador cadastrado no catálogo.", "success")
        return redirect(url_for("team_indicador_catalogo"))

    @app.route("/catalogo/indicador/<int:did>", methods=["POST"])
    @controladoria_required
    def team_indicador_def_edit(did):
        from team.models import IndicatorDef
        d = db.session.get(IndicatorDef, did) or abort(404)
        _def_from_form(d, request.form)
        # propaga nome/dimensão/tipo para os usos (mantém exibição consistente)
        for ind in Indicator.query.filter_by(indicator_def_id=d.id).all():
            ind.title = d.title
            ind.dimension = d.dimension
            ind.target_type = d.target_type
        db.session.commit()
        flash("Indicador do catálogo atualizado.", "success")
        return redirect(url_for("team_indicador_catalogo"))

    @app.route("/catalogo/indicador/<int:did>/toggle", methods=["POST"])
    @controladoria_required
    def team_indicador_def_toggle(did):
        from team.models import IndicatorDef
        d = db.session.get(IndicatorDef, did) or abort(404)
        d.active = not d.active
        db.session.commit()
        return redirect(url_for("team_indicador_catalogo"))

    @app.route("/catalogo/indicador/<int:did>/excluir", methods=["POST"])
    @controladoria_required
    def team_indicador_def_delete(did):
        from team.models import IndicatorDef
        d = db.session.get(IndicatorDef, did) or abort(404)
        if d.usos:
            flash(f"Não dá para excluir: em uso em {d.usos} painel(éis). "
                  "Desative-o se quiser tirá-lo da lista.", "warning")
            return redirect(url_for("team_indicador_catalogo"))
        db.session.delete(d)
        db.session.commit()
        flash("Indicador removido do catálogo.", "success")
        return redirect(url_for("team_indicador_catalogo"))

    # ---------------- Medição única: um indicador do catálogo, todos os painéis --
    def _usos_ativos(did):
        """Indicadores (um por painel) que usam a definição, só painéis ativos."""
        from team.models_workflow import Panel
        inds = (Indicator.query.filter_by(indicator_def_id=did, active=True)
                .outerjoin(Panel, Indicator.panel_id == Panel.id)
                .filter(db.or_(Panel.id.is_(None), Panel.active.is_(True)))
                .order_by(Panel.sort_order, Panel.name).all())
        return [i for i in inds if not (i.member and not i.member.active)]

    @app.route("/indicadores/medir", methods=["GET", "POST"])
    @controladoria_required
    def team_indicador_medir():
        from team.models import IndicatorDef
        if request.method == "POST":
            return _grava_medicao()
        defs = (IndicatorDef.query.filter_by(active=True)
                .order_by(IndicatorDef.sort_order, IndicatorDef.title).all())
        usos = {d.id: _usos_ativos(d.id) for d in defs}
        defs = [d for d in defs if usos[d.id]]
        # os comuns a vários painéis primeiro
        defs.sort(key=lambda d: (-len(usos[d.id]), d.sort_order or 100, d.title))
        sel = db.session.get(IndicatorDef, request.args.get("def_id", type=int) or 0)
        comp = _current_competency()
        return render_template("team/indicador_medir.html", defs=defs, usos=usos,
                               sel=sel, inds=usos.get(sel.id, []) if sel else [],
                               comps=_comps(), comp_atual=comp)

    def _grava_medicao():
        from team.models import IndicatorDef
        import shutil
        d = db.session.get(IndicatorDef, _int(request.form.get("def_id"))) or abort(404)
        comp_id = _int(request.form.get("competency_id"))
        comp = db.session.get(Competency, comp_id) if comp_id else None
        periodo = (request.form.get("period_label") or "").strip() or (comp.label if comp else None)
        if not comp and not periodo:
            flash("Informe a competência ou o período da medição.", "danger")
            return redirect(url_for("team_indicador_medir", def_id=d.id))
        geral = request.form.get("outcome") or "na"
        valor = (request.form.get("value_label") or "").strip() or None
        nota = (request.form.get("note") or "").strip() or None
        file = request.files.get("evidence")
        if file and file.filename:
            ext = os.path.splitext(file.filename)[1].lower()
            if ext not in ALLOWED_EVIDENCE:
                flash("Formato de evidência não permitido.", "danger")
                return redirect(url_for("team_indicador_medir", def_id=d.id))
        marcados = {int(x) for x in request.form.getlist("ind") if x.isdigit()}
        alvo = [i for i in _usos_ativos(d.id) if i.id in marcados]
        if not alvo:
            flash("Marque ao menos um painel.", "warning")
            return redirect(url_for("team_indicador_medir", def_id=d.id))

        agora = datetime.utcnow()
        origem = None                     # 1º arquivo salvo; os demais são cópias
        novos = atualizados = 0
        for i in alvo:
            q = IndicatorResult.query.filter_by(indicator_id=i.id)
            q = (q.filter_by(competency_id=comp.id) if comp
                 else q.filter_by(competency_id=None, period_label=periodo))
            r = q.first()
            if r:
                atualizados += 1
            else:
                r = IndicatorResult(indicator_id=i.id)
                db.session.add(r)
                novos += 1
            r.competency_id = comp.id if comp else None
            r.period_label = periodo
            r.outcome = request.form.get(f"outcome_{i.id}") or geral
            r.value_label = valor
            r.note = nota
            r.computed = False
            r.recorded_by = current_user.id
            r.recorded_at = agora
            if file and file.filename:
                # cada resultado tem a sua cópia: excluir num painel não apaga nos outros
                folder = os.path.join(EVIDENCE_DIR, str(i.id))
                os.makedirs(folder, exist_ok=True)
                path = os.path.join(folder, secure_filename(
                    f"{agora:%Y%m%d%H%M%S}_{file.filename}"))
                if origem is None:
                    file.save(path)
                    origem = path
                else:
                    shutil.copyfile(origem, path)
                if r.evidence_path and r.evidence_path != path and os.path.exists(r.evidence_path):
                    try:
                        os.remove(r.evidence_path)
                    except OSError:
                        pass
                r.evidence_path = path
                r.evidence_name = file.filename
                r.evidence_by = current_user.id
                r.evidence_at = agora
        db.session.commit()
        log_audit(current_user.id, "indicador_medicao_unica", "indicator_def",
                  f"{d.id}: {len(alvo)} painel(éis), {periodo}")
        partes = []
        if novos:
            partes.append(f"{novos} novo(s)")
        if atualizados:
            partes.append(f"{atualizados} atualizado(s) — já havia registro em {periodo}")
        flash(f"“{d.title}” medido em {len(alvo)} painel(éis): " + ", ".join(partes) + ".",
              "success")
        return redirect(url_for("team_indicador_medir", def_id=d.id))

    @app.route("/indicador/<int:iid>")
    @team_required
    def team_indicador(iid):
        i = db.session.get(Indicator, iid) or abort(404)
        results = (IndicatorResult.query.filter_by(indicator_id=i.id)
                   .order_by(IndicatorResult.recorded_at.desc()).all())
        outros = (len([u for u in _usos_ativos(i.indicator_def_id) if u.id != i.id])
                  if i.indicator_def_id else 0)
        return render_template("team/indicador_detail.html", i=i, results=results,
                               outros_paineis=outros,
                               mm=_member_map(), comps=_comps(),
                               users={u.id: u for u in User.query.all()})

    @app.route("/indicador/<int:iid>/resultado", methods=["POST"])
    @team_required
    def team_indicador_result(iid):
        i = db.session.get(Indicator, iid) or abort(404)
        r = IndicatorResult(
            indicator_id=i.id,
            competency_id=_int(request.form.get("competency_id")),
            period_label=request.form.get("period_label") or None,
            outcome=request.form.get("outcome") or "na",
            value_label=request.form.get("value_label") or None,
            note=request.form.get("note") or None,
            recorded_by=current_user.id)
        # evidencia (anexo simples, com quem/quando)
        file = request.files.get("evidence")
        if file and file.filename:
            ext = os.path.splitext(file.filename)[1].lower()
            if ext not in ALLOWED_EVIDENCE:
                flash("Formato de evidência não permitido.", "danger")
                return redirect(url_for("team_indicador", iid=i.id))
            folder = os.path.join(EVIDENCE_DIR, str(i.id))
            os.makedirs(folder, exist_ok=True)
            fname = secure_filename(f"{datetime.utcnow():%Y%m%d%H%M%S}_{file.filename}")
            path = os.path.join(folder, fname)
            file.save(path)
            r.evidence_path = path
            r.evidence_name = file.filename
            r.evidence_by = current_user.id
            r.evidence_at = datetime.utcnow()
        db.session.add(r)
        db.session.commit()
        log_audit(current_user.id, "indicador_resultado", "indicator", str(i.id))
        flash("Resultado registrado.", "success")
        return redirect(url_for("team_indicador", iid=i.id))

    @app.route("/indicador/resultado/<int:rid>/evidencia")
    @team_required
    def team_evidence_download(rid):
        r = db.session.get(IndicatorResult, rid) or abort(404)
        if not r.evidence_path or not os.path.exists(r.evidence_path):
            abort(404)
        return send_file(r.evidence_path, as_attachment=True,
                         download_name=r.evidence_name or "evidencia")

    @app.route("/indicador/resultado/<int:rid>/excluir", methods=["POST"])
    @team_required
    def team_result_delete(rid):
        r = db.session.get(IndicatorResult, rid) or abort(404)
        iid = r.indicator_id
        # remove arquivo fisico da evidencia, se houver
        if r.evidence_path and os.path.exists(r.evidence_path):
            try:
                os.remove(r.evidence_path)
            except OSError:
                pass
        db.session.delete(r)
        db.session.commit()
        flash("Resultado removido.", "success")
        return redirect(url_for("team_indicador", iid=iid))

    @app.route("/indicador/<int:iid>/auto", methods=["POST"])
    @team_required
    def team_indicador_auto(iid):
        """Calcula automaticamente o atingimento a partir do motor de fechamento."""
        i = db.session.get(Indicator, iid) or abort(404)
        comp = _current_competency()
        n = _autocompute_indicator(i, comp)
        if n:
            flash(f"Atingimento calculado a partir do fechamento ({n} período[s]).",
                  "success")
        else:
            flash("Sem dados de fechamento suficientes para cálculo automático.",
                  "warning")
        return redirect(url_for("team_indicador", iid=i.id))

    def _autocompute_indicator(i, comp):
        """Deriva o resultado de metas de prazo (D+x) a partir das atividades
        auto_metric concluidas da competencia. A submissao/atividade e a evidencia."""
        if not comp or i.target_type != "prazo_du":
            return 0
        # procura atividade de fechamento do mesmo membro concluida na competencia
        acts = (Activity.query.filter_by(competency_id=comp.id, member_id=i.member_id)
                .filter(Activity.kind.in_(["fechamento", "recorrente"]))
                .filter(Activity.status == "concluida").all())
        if not acts:
            return 0
        # ja existe resultado desta competencia?
        exists = IndicatorResult.query.filter_by(
            indicator_id=i.id, competency_id=comp.id).first()
        if exists:
            return 0
        # avalia on-time: concluida ate o prazo?
        on_time = all((a.done_at.date() <= a.due_date) if (a.done_at and a.due_date)
                      else True for a in acts)
        r = IndicatorResult(
            indicator_id=i.id, competency_id=comp.id, period_label=comp.label,
            outcome="atingido" if on_time else "nao_atingido",
            value_label=f"{i.target_label} — {'no prazo' if on_time else 'com atraso'}",
            computed=True, note="Calculado automaticamente pelo motor de fechamento.",
            recorded_by=current_user.id)
        db.session.add(r)
        db.session.commit()
        return 1

    # ==================================================================
    # PROJETOS
    # ==================================================================
    @app.route("/projetos")
    @team_required
    def team_projetos():
        pq = Project.query
        sid = _scoped_member_id()
        if sid is not None:               # profissional: projetos que ele lidera
            pq = pq.filter(Project.owner_member_id == sid)
        projects = pq.order_by(Project.status, Project.sort_order,
                               Project.name).all()
        # progresso por marcos
        prog = {}
        for p in projects:
            ms = p.milestones
            done = sum(1 for m in ms if m.done)
            prog[p.id] = {"done": done, "total": len(ms),
                          "pct": round(done / len(ms) * 100) if ms else 0,
                          "open_acts": Activity.query.filter_by(project_id=p.id)
                          .filter(Activity.status.in_(["pendente", "em_andamento",
                                                       "bloqueada"])).count()}
        return render_template("team/projetos.html", projects=projects, prog=prog,
                               members=_members(), mm=_member_map(),
                               gantt=build_gantt(projects, prog),
                               today=fuso.hoje())

    @app.route("/projeto/novo", methods=["POST"])
    @team_required
    def team_projeto_new():
        p = Project(
            name=(request.form.get("name") or "").strip() or "Novo projeto",
            code=request.form.get("code") or None,
            status=request.form.get("status") or "ativo",
            owner_member_id=_int(request.form.get("owner_member_id")),
            manager_member_id=_int(request.form.get("manager_member_id")),
            description=request.form.get("description") or None,
            start_date=_date(request.form.get("start_date")),
            target_date=_date(request.form.get("target_date")),
            confidential=bool(request.form.get("confidential")))
        db.session.add(p)
        db.session.commit()
        flash("Projeto criado.", "success")
        return redirect(url_for("team_projeto", pid=p.id))

    @app.route("/projeto/<int:pid>")
    @team_required
    def team_projeto(pid):
        p = db.session.get(Project, pid) or abort(404)
        acts = (Activity.query.filter_by(project_id=p.id)
                .order_by(Activity.status, Activity.due_date).all())
        timeline = build_activity_timeline(acts, color=p.color or "#1d5da8")
        from team.models_workflow import DeadlineRevision
        mids = [m.id for m in p.milestones]
        aids = [a.id for a in acts]
        escopo = db.or_(
            db.and_(DeadlineRevision.entity_type == "project",
                    DeadlineRevision.entity_id == p.id),
            db.and_(DeadlineRevision.entity_type == "milestone",
                    DeadlineRevision.entity_id.in_(mids or [-1])),
            db.and_(DeadlineRevision.entity_type == "activity",
                    DeadlineRevision.entity_id.in_(aids or [-1])))
        todas = (DeadlineRevision.query.filter(escopo)
                 .order_by(DeadlineRevision.created_at.desc()).all())
        revisoes = [r for r in todas if r.status == "aplicada"]
        pendentes = [r for r in todas if r.status == "solicitada"]
        return render_template("team/projeto_detail.html", p=p, acts=acts,
                               members=_members(), mm=_member_map(),
                               timeline=timeline, revisoes=revisoes,
                               pendentes=pendentes, pode_gerir=_pode_gerir_projeto(p),
                               lider=_projeto_lider(p), lider_ferias=_lider_em_ferias(p),
                               gestor=_projeto_gestor(p),
                               gestor_ferias=_membro_em_ferias(_projeto_gestor(p)),
                               users={u.id: u for u in User.query.all()},
                               today=fuso.hoje())

    @app.route("/projeto/<int:pid>/editar", methods=["POST"])
    @team_required
    def team_projeto_edit(pid):
        """Edita os dados do projeto — em especial as datas que alimentam o Gantt."""
        p = db.session.get(Project, pid) or abort(404)
        p.name = (request.form.get("name") or p.name).strip()
        p.code = (request.form.get("code") or "").strip() or None
        p.owner_member_id = _int(request.form.get("owner_member_id"))
        p.manager_member_id = _int(request.form.get("manager_member_id"))
        p.start_date = _date(request.form.get("start_date"))
        p.target_date = _date(request.form.get("target_date"))
        p.color = request.form.get("color") or p.color
        p.description = (request.form.get("description") or "").strip() or None
        p.confidential = bool(request.form.get("confidential"))
        db.session.commit()
        log_audit(current_user.id, "projeto_editado", "project", p.name)
        flash("Projeto atualizado.", "success")
        return redirect(url_for("team_projeto", pid=p.id))

    @app.route("/marco/<int:mid>/editar", methods=["POST"])
    @team_required
    def team_milestone_edit(mid):
        m = db.session.get(Milestone, mid) or abort(404)
        m.title = (request.form.get("title") or m.title).strip()
        m.due_date = _date(request.form.get("due_date"))
        db.session.commit()
        return redirect(url_for("team_projeto", pid=m.project_id))

    @app.route("/marco/<int:mid>/excluir", methods=["POST"])
    @team_required
    def team_milestone_delete(mid):
        m = db.session.get(Milestone, mid) or abort(404)
        pid = m.project_id
        db.session.delete(m)
        db.session.commit()
        flash("Marco removido.", "success")
        return redirect(url_for("team_projeto", pid=pid))

    @app.route("/projeto/<int:pid>/marco", methods=["POST"])
    @team_required
    def team_milestone_new(pid):
        p = db.session.get(Project, pid) or abort(404)
        m = Milestone(project_id=p.id,
                      title=(request.form.get("title") or "").strip() or "Marco",
                      due_date=_date(request.form.get("due_date")),
                      sort_order=_int(request.form.get("sort_order")) or 100)
        db.session.add(m)
        db.session.commit()
        return redirect(url_for("team_projeto", pid=p.id))

    @app.route("/marco/<int:mid>/toggle", methods=["POST"])
    @team_required
    def team_milestone_toggle(mid):
        m = db.session.get(Milestone, mid) or abort(404)
        m.done = not m.done
        m.done_at = datetime.utcnow() if m.done else None
        db.session.commit()
        return redirect(url_for("team_projeto", pid=m.project_id))

    @app.route("/projeto/<int:pid>/status", methods=["POST"])
    @team_required
    def team_projeto_status(pid):
        p = db.session.get(Project, pid) or abort(404)
        st = request.form.get("status")
        if st in ("planejado", "ativo", "pausado", "concluido", "cancelado"):
            p.status = st
            db.session.commit()
        return redirect(request.referrer or url_for("team_projetos"))

    @app.route("/projeto/<int:pid>/confidencial", methods=["POST"])
    @team_required
    def team_projeto_confidencial(pid):
        """Bloqueia/desbloqueia o projeto (confidencial: nome não vai a canais externos)."""
        p = db.session.get(Project, pid) or abort(404)
        p.confidential = not p.confidential
        db.session.commit()
        flash("Projeto bloqueado (confidencial)." if p.confidential
              else "Projeto desbloqueado.", "success")
        return redirect(request.referrer or url_for("team_projeto", pid=p.id))

    @app.route("/projeto/<int:pid>/abrir", methods=["POST"])
    @team_required
    def team_projeto_abrir(pid):
        """Abertura: marca como ativo e registra a data de início se faltar."""
        p = db.session.get(Project, pid) or abort(404)
        p.status = "ativo"
        if not p.start_date:
            p.start_date = fuso.hoje()
        db.session.commit()
        flash("Projeto aberto.", "success")
        return redirect(request.referrer or url_for("team_projeto", pid=p.id))

    @app.route("/projeto/<int:pid>/encerrar", methods=["POST"])
    @team_required
    def team_projeto_encerrar(pid):
        """Encerramento: conclui o projeto e registra a data de término se faltar."""
        p = db.session.get(Project, pid) or abort(404)
        p.status = "concluido"
        if not p.target_date:
            p.target_date = fuso.hoje()
        db.session.commit()
        flash("Projeto encerrado.", "success")
        return redirect(request.referrer or url_for("team_projeto", pid=p.id))

    @app.route("/projeto/<int:pid>/revisar-prazo", methods=["POST"])
    @team_required
    def team_projeto_revisar_prazo(pid):
        """Redefinir o prazo do PROJETO: só líder do projeto ou gestor."""
        p = db.session.get(Project, pid) or abort(404)
        if not _pode_gerir_projeto(p):
            flash("Só o líder do projeto ou o gestor podem redefinir o prazo do projeto.",
                  "danger")
            return redirect(url_for("team_projeto", pid=p.id))
        log_audit(current_user.id, "prazo_revisado", "project", p.name)
        return _revisar_prazo("project", p.id, p.target_date,
                              f"Projeto {p.name}", url_for("team_projeto", pid=p.id))

    @app.route("/atividade/<int:aid>/revisar-prazo", methods=["POST"])
    @team_required
    def team_atividade_revisar_prazo(aid):
        """Realinhar prazo de atividade: líder/gestor aplicam; demais solicitam."""
        a = db.session.get(Activity, aid) or abort(404)
        return _revisar_prazo("activity", a.id, a.due_date,
                              f"Atividade {a.title}", url_for("team_atividade", aid=a.id))

    @app.route("/marco/<int:mid>/revisar-prazo", methods=["POST"])
    @team_required
    def team_milestone_revisar_prazo(mid):
        m = db.session.get(Milestone, mid) or abort(404)
        return _revisar_prazo("milestone", m.id, m.due_date,
                              f"Marco {m.title}", url_for("team_projeto", pid=m.project_id))

    @app.route("/revisao/<int:rid>/aprovar", methods=["POST"])
    @team_required
    def team_revisao_aprovar(rid):
        from team.models_workflow import DeadlineRevision
        from models import notify
        rev = db.session.get(DeadlineRevision, rid) or abort(404)
        if rev.status != "solicitada" or not _pode_aprovar_revisao(
                rev.entity_type, rev.entity_id):
            abort(403)
        rev.status = "aplicada"
        rev.approved_by = current_user.id
        rev.approved_at = datetime.utcnow()
        _aplicar_revisao(rev)
        db.session.commit()
        if rev.requested_by:
            notify(rev.requested_by, "Realinhamento de prazo aprovado",
                   f"Novo prazo: {rev.new_date.strftime('%d/%m/%Y')}.",
                   kind="prazo", url=request.referrer or url_for("team_projetos"))
        flash("Realinhamento aprovado e aplicado.", "success")
        return redirect(request.referrer or url_for("team_projetos"))

    @app.route("/revisao/<int:rid>/recusar", methods=["POST"])
    @team_required
    def team_revisao_recusar(rid):
        from team.models_workflow import DeadlineRevision
        from models import notify
        rev = db.session.get(DeadlineRevision, rid) or abort(404)
        if rev.status != "solicitada" or not _pode_aprovar_revisao(
                rev.entity_type, rev.entity_id):
            abort(403)
        rev.status = "recusada"
        rev.approved_by = current_user.id
        rev.approved_at = datetime.utcnow()
        rev.decision_note = (request.form.get("decision_note") or "").strip() or None
        db.session.commit()
        if rev.requested_by:
            notify(rev.requested_by, "Realinhamento de prazo recusado",
                   (rev.decision_note or "Sem justificativa."),
                   kind="prazo", url=request.referrer or url_for("team_projetos"))
        flash("Realinhamento recusado.", "info")
        return redirect(request.referrer or url_for("team_projetos"))

    # ==================================================================
    # CAPACIDADE — mapa de calor pessoa x dia util
    # ==================================================================
    @app.route("/capacidade")
    @controladoria_required
    def team_capacidade():
        """Onde a carga se concentra no calendario do fechamento.

        Responde 'quem vai estourar e em que dia' ANTES do atraso acontecer.
        """
        from engine.calendar_br import business_days
        hoje = fuso.hoje()
        comp = _current_competency()
        # janela: mes do prazo da competencia aberta (onde acontece o pico)
        ref = comp.deadline if (comp and comp.deadline) else hoje
        ano, mes = ref.year, ref.month
        if request.args.get("mes"):
            try:
                ano, mes = [int(x) for x in request.args["mes"].split("-")]
            except (ValueError, IndexError):
                pass
        dias = business_days(ano, mes)
        membros = _members()

        # atividades em aberto por (membro, dia) + atrasadas e sem prazo
        grade = {}          # (member_id, iso) -> [atividades]
        atrasadas = {}      # member_id -> [atividades]
        sem_prazo = {}      # member_id -> [atividades]
        q = Activity.query.filter(Activity.status.in_(
            ["pendente", "em_andamento", "bloqueada"]))
        _sid = _scoped_member_id()
        if _sid is not None:
            q = q.filter(Activity.member_id == _sid)
        for a in q.all():
            mid = a.member_id or 0
            if not a.due_date or a.due_provisional:
                sem_prazo.setdefault(mid, []).append(a)
                continue
            if a.due_date < hoje:
                atrasadas.setdefault(mid, []).append(a)
                continue
            if a.due_date.year == ano and a.due_date.month == mes:
                grade.setdefault((mid, a.due_date.isoformat()), []).append(a)

        # pico do dia (para calibrar a intensidade da cor)
        por_dia = {}
        for (mid, iso), lst in grade.items():
            por_dia[iso] = por_dia.get(iso, 0) + len(lst)
        pico = max(por_dia.values()) if por_dia else 0

        # ausencias: um dia marcado como fora nao deveria receber carga
        from team.models_workflow import absences_map
        fora = absences_map(dias)

        linhas = []
        for m in membros:
            celulas = []
            ausente = fora.get(m.id, {})
            for d in dias:
                lst = grade.get((m.id, d.isoformat()), [])
                aus = ausente.get(d.isoformat())
                celulas.append({"dia": d, "n": len(lst), "acts": lst[:6],
                                "hoje": d == hoje,
                                "ausente": aus.kind_pt if aus else None})
            # conflito: atividade marcada para um dia em que a pessoa esta fora
            conflitos = sum(c["n"] for c in celulas if c["ausente"])
            linhas.append({
                "m": m, "celulas": celulas,
                "atrasadas": len(atrasadas.get(m.id, [])),
                "sem_prazo": len(sem_prazo.get(m.id, [])),
                "total": sum(c["n"] for c in celulas),
                "conflitos": conflitos,
                "dias_fora": len(ausente),
            })
        # linha "sem responsavel"
        orfas = {"atrasadas": len(atrasadas.get(0, [])),
                 "sem_prazo": len(sem_prazo.get(0, [])),
                 "celulas": [{"dia": d, "n": len(grade.get((0, d.isoformat()), [])),
                              "acts": grade.get((0, d.isoformat()), [])[:6],
                              "hoje": d == hoje} for d in dias]}
        orfas["total"] = sum(c["n"] for c in orfas["celulas"])

        totais_dia = [{"dia": d, "n": por_dia.get(d.isoformat(), 0),
                       "hoje": d == hoje} for d in dias]
        return render_template("team/capacidade.html", dias=dias, linhas=linhas,
                               orfas=orfas, totais_dia=totais_dia, pico=pico,
                               comp=comp, ano=ano, mes=mes, hoje=hoje,
                               deadline=(comp.deadline if comp else None))

    # ==================================================================
    # CARTEIRA (empresa x pessoa) — o "pilotar paineis"
    # ==================================================================
    @app.route("/carteira")
    @team_required
    def team_carteira():
        f_seat = request.args.get("seat", type=int)
        f_member = request.args.get("member_id", type=int)
        f_seg = request.args.get("segment")
        q = (CompanyAssignment.query.join(Company)
             .order_by(CompanyAssignment.member_id, Company.name))
        sid = _scoped_member_id()
        if sid is not None:               # profissional: so a propria carteira
            q = q.filter(CompanyAssignment.member_id == sid)
        if f_seat:
            q = q.filter(CompanyAssignment.seat == f_seat)
        if f_member and sid is None:
            q = q.filter(CompanyAssignment.member_id == f_member)
        if f_seg:
            q = q.filter(CompanyAssignment.segment == f_seg)
        rows = q.all()
        segments = sorted({a.segment for a in CompanyAssignment.query.all() if a.segment})
        # carga por membro (escopada quando profissional)
        load = {}
        cq = CompanyAssignment.query
        if sid is not None:
            cq = cq.filter(CompanyAssignment.member_id == sid)
        for a in cq.all():
            if a.member_id:
                d = load.setdefault(a.member_id, {"n": 0, "real": 0, "ideal": 0})
                d["n"] += 1
                d["real"] += a.load_real or 0
                d["ideal"] += a.load_ideal or 0
        return render_template("team/carteira.html", rows=rows, members=_members(),
                               mm=_member_map(), segments=segments,
                               f_member=f_member, f_seg=f_seg, load=load,
                               escopo=(sid is not None))

    @app.route("/carteira/<int:aid>/atribuir", methods=["POST"])
    @team_required
    def team_carteira_assign(aid):
        a = db.session.get(CompanyAssignment, aid) or abort(404)
        a.member_id = _int(request.form.get("member_id"))
        db.session.commit()
        if request.form.get("ajax"):
            m = db.session.get(TeamMember, a.member_id) if a.member_id else None
            return jsonify(ok=True, member=m.name if m else "—")
        flash(f"{a.company.name} atribuída.", "success")
        return redirect(request.referrer or url_for("team_carteira"))

    @app.route("/carteira/atribuir-lote", methods=["POST"])
    @team_required
    def team_carteira_bulk():
        """Atribui todas as empresas de um assento a um membro (pilotar rapido)."""
        seat = _int(request.form.get("seat"))
        member_id = _int(request.form.get("member_id"))
        if seat:
            n = CompanyAssignment.query.filter_by(seat=seat).update(
                {"member_id": member_id})
            db.session.commit()
            log_audit(current_user.id, "carteira_lote", "assignment",
                      f"seat={seat} -> member={member_id} ({n})")
            flash(f"{n} empresa(s) do assento {seat} atribuída(s).", "success")
        return redirect(url_for("team_carteira"))

    # ==================================================================
    # ALERTAS (matriz evento x canal)
    # ==================================================================
    @app.route("/alertas")
    @team_required
    def team_alertas():
        ensure_alert_defaults()
        settings = {s.event_key: s for s in AlertChannelSetting.query.all()}
        ordered = [(k, lbl, settings.get(k)) for k, lbl in ALERT_EVENTS]
        recent = (AlertLog.query.order_by(AlertLog.created_at.desc()).limit(40).all())
        # contadores do log
        stats = {}
        for lg in AlertLog.query.all():
            stats[lg.status] = stats.get(lg.status, 0) + 1
        me = _current_member()
        from team.models_workflow import PushSubscription
        n_aparelhos = PushSubscription.query.filter_by(user_id=current_user.id).count()
        return render_template("team/alertas.html", ordered=ordered, recent=recent,
                               n_aparelhos=n_aparelhos,
                               mm=_member_map(), stats=stats,
                               email_on=team_alerts.EMAIL_ENABLED,
                               whatsapp_on=team_alerts.WHATSAPP_ENABLED,
                               email_via=team_alerts.email_via(),
                               meu_email=current_user.email,
                               meu_whatsapp=(me.whatsapp if me else None))

    @app.route("/alertas/teste", methods=["POST"])
    @team_required
    def team_alertas_teste():
        """Envia um e-mail e um WhatsApp de teste para quem clicou.

        Serve para conferir as credenciais (Gmail/Twilio) sem esperar o ciclo."""
        canal = request.form.get("canal")
        digitado = (request.form.get("destino") or "").strip()
        if canal == "email":
            destino = digitado or current_user.email
            if "@" not in (destino or ""):
                flash("Informe um e-mail válido para o teste.", "warning")
                return redirect(url_for("team_alertas"))
            ok, err = team_alerts.send_email(
                destino, "Teste de alerta — Portal Controladoria",
                "Se você recebeu este e-mail, o canal de e-mail do portal está "
                f"funcionando.\n\nPortal: {team_alerts.PORTAL_URL}")
        else:
            me = _current_member()
            numero = team_alerts.normaliza_whatsapp(
                digitado or (me.whatsapp if me else None))
            if not numero:
                flash("Informe o número de WhatsApp para o teste (ou cadastre o seu "
                      "em Administração › Time com o seu login vinculado).", "warning")
                return redirect(url_for("team_alertas"))
            ok, err = team_alerts.send_whatsapp(
                numero, "Controladoria J&F: teste de alerta. O canal de WhatsApp "
                        "do portal está funcionando.")
            destino = numero
        log_audit(current_user.id, f"alerta_teste_{canal}", "team",
                  f"{destino} ok={ok} {err or ''}")
        if ok:
            flash(f"Teste enviado para {destino}. Confira a caixa/WhatsApp.", "success")
        else:
            flash(f"Não foi possível enviar para {destino}: {_motivo(err)}", "danger")
        return redirect(url_for("team_alertas"))

    def _motivo(err):
        return {
            "email_desligado": "canal de e-mail desligado — faltam SMTP_USER/SMTP_PASSWORD no servidor.",
            "smtp_sem_senha": "falta a variável SMTP_PASSWORD (senha de app do Gmail).",
            "sem_destinatario": "seu usuário não tem e-mail.",
            "whatsapp_desligado": "canal de WhatsApp desligado — faltam as credenciais TWILIO_* no servidor.",
            "credenciais_twilio_ausentes": "faltam TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN ou TWILIO_WHATSAPP_FROM.",
        }.get(err, err or "erro desconhecido")

    @app.route("/alertas/salvar", methods=["POST"])
    @team_required
    def team_alertas_save():
        ensure_alert_defaults()
        for key, _lbl in ALERT_EVENTS:
            s = db.session.get(AlertChannelSetting, key)
            if not s:
                continue
            s.email = bool(request.form.get(f"{key}__email"))
            s.whatsapp = bool(request.form.get(f"{key}__whatsapp"))
            s.painel = bool(request.form.get(f"{key}__painel"))
            s.push = bool(request.form.get(f"{key}__push"))
            s.escalate_manager = bool(request.form.get(f"{key}__escalate"))
            ld = request.form.get(f"{key}__lead")
            if ld is not None and str(ld).strip():
                try:
                    s.lead_days = int(ld)
                except ValueError:
                    pass
        db.session.commit()
        log_audit(current_user.id, "alertas_config", "team", "matriz salva")
        flash("Matriz de alertas salva.", "success")
        return redirect(url_for("team_alertas"))

    @app.route("/alertas/rodar", methods=["POST"])
    @team_required
    def team_alertas_run():
        dry = bool(request.form.get("dry_run"))
        summary = team_alerts.run_alert_cycle(dry_run=dry)
        canais = ", ".join(f"{k}={v}" for k, v in summary["por_canal"].items()) or "nenhum"
        flash(f"Ciclo {'(simulação) ' if dry else ''}concluído — "
              f"lembretes: {summary['lembrete_previo']}, vencem hoje: "
              f"{summary['vence_hoje']}, atrasos: {summary['atraso']}. "
              f"Canais: {canais}.", "success")
        return redirect(url_for("team_alertas"))

    # ------------------------------------------------------------------
    # Helpers de formulario
    # ------------------------------------------------------------------
    def _companies():
        return Company.query.filter_by(active=True).order_by(Company.name).all()

    # ==================================================================
    # EXPORTAÇÕES PARA EXCEL (botões "Excel" das telas) — respeitam o escopo
    # do perfil: profissional exporta só o que vê na tela.
    # ==================================================================
    def _planilha(nome, titulo, cabecalho, linhas, larguras=None):
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = titulo[:31]
        ws.append(cabecalho)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="14273A")
            c.alignment = Alignment(vertical="center")
        for ln in linhas:
            ws.append(list(ln))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, _ in enumerate(cabecalho, 1):
            w = (larguras or {}).get(i)
            if not w:
                w = min(60, max([len(str(cabecalho[i - 1]))] +
                                [len(str(r[i - 1])) for r in linhas if r[i - 1] is not None] or [10]) + 2)
            ws.column_dimensions[get_column_letter(i)].width = w
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        log_audit(current_user.id, "exportou_excel", "export", nome)
        return send_file(buf, as_attachment=True,
                         download_name=f"{nome}_{fuso.hoje():%Y-%m-%d}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def _d(v):
        return v.strftime("%d/%m/%Y") if v else None

    @app.route("/export/indicadores.xlsx")
    @team_required
    def team_export_indicadores():
        from team.models_workflow import Panel
        q = Indicator.query.filter_by(active=True)
        sid = _scoped_member_id()
        if sid is not None:
            me = _current_member()
            ids = {me.panel_id} if me and me.panel_id else set()
            proprio = Panel.query.filter_by(kind="pessoal", owner_member_id=me.id).first() if me else None
            if proprio:
                ids.add(proprio.id)
            q = q.filter(Indicator.panel_id.in_(ids or {-1}))
        pmap = {p.id: p.name for p in Panel.query.all()}
        linhas = []
        for i in q.order_by(Indicator.panel_id, Indicator.seq).all():
            ult = max(i.results, key=lambda r: r.recorded_at, default=None) if i.results else None
            linhas.append([pmap.get(i.panel_id, ""), i.seq, i.title, i.dimension, i.target_label,
                           i.unit, {"menor": "menor melhor", "maior": "maior melhor"}.get(i.direction or "", ""),
                           i.weight, i.scale_min, i.scale_obj, i.scale_sup,
                           (ult.outcome if ult else None), (ult.period_label if ult else None)])
        return _planilha("indicadores", "Indicadores",
                         ["Painel", "Nº", "Meta", "Dimensão", "Indicador", "Unidade", "Sentido",
                          "Peso (%)", "Mínimo", "Objetivo", "Superado", "Último resultado", "Período"], linhas)

    @app.route("/export/atividades.xlsx")
    @team_required
    def team_export_atividades():
        q = Activity.query
        sid = _scoped_member_id()
        if sid is not None:
            q = q.filter(Activity.member_id == sid)
        cid = request.args.get("competency_id", type=int)
        if cid:
            q = q.filter(Activity.competency_id == cid)
        ref = fuso.hoje()
        st = {"atrasada": "Atrasada", "vence_hoje": "Vence hoje", "pendente": "Pendente",
              "em_andamento": "Em andamento", "aguardando": "Aguardando insumo",
              "bloqueada": "Bloqueada", "concluida": "Concluída", "cancelada": "Cancelada"}
        linhas = [[a.title, a.kind_pt(), (a.member.name if a.member else None),
                   (a.company.name if a.company else None),
                   (a.competency.label if a.competency else None), _d(a.due_date),
                   st.get(a.effective_status(ref), a.effective_status(ref)), a.priority,
                   (fuso.local(a.done_at).strftime("%d/%m/%Y %H:%M") if a.done_at else None)]
                  for a in q.order_by(Activity.due_date.is_(None), Activity.due_date, Activity.id).all()]
        return _planilha("atividades", "Atividades",
                         ["Atividade", "Tipo", "Responsável", "Empresa", "Competência", "Prazo",
                          "Situação", "Prioridade", "Concluída em"], linhas)

    @app.route("/export/carteira.xlsx")
    @team_required
    def team_export_carteira():
        q = CompanyAssignment.query.join(Company).order_by(Company.name)
        sid = _scoped_member_id()
        if sid is not None:
            q = q.filter(CompanyAssignment.member_id == sid)
        linhas = [[a.company.name, a.company.code, a.segment, (a.member.name if a.member else None),
                   a.flow, a.responsibility, ", ".join(a.deliverables), a.load_real, a.load_ideal,
                   a.production, a.note] for a in q.all()]
        return _planilha("carteira", "Carteira",
                         ["Empresa", "Código", "Segmento", "Responsável", "Fluxo", "Responsabilidade",
                          "Entregas", "Suportes", "Suporte ideal", "Produção", "Observação"], linhas)

    @app.route("/export/projetos.xlsx")
    @team_required
    def team_export_projetos():
        pq = Project.query
        sid = _scoped_member_id()
        if sid is not None:
            pq = pq.filter(Project.owner_member_id == sid)
        linhas = []
        for p in pq.order_by(Project.status, Project.sort_order, Project.name).all():
            ms = p.milestones
            feitos = sum(1 for m in ms if m.done)
            linhas.append([p.name, p.code, p.status_pt, (p.owner.name if p.owner else None),
                           (p.manager.name if p.manager else None), _d(p.start_date), _d(p.target_date),
                           f"{feitos}/{len(ms)}" if ms else "0/0",
                           Activity.query.filter_by(project_id=p.id)
                           .filter(Activity.status.in_(["pendente", "em_andamento", "bloqueada"])).count()])
        return _planilha("projetos", "Projetos",
                         ["Projeto", "Código", "Situação", "Líder", "Gestor", "Início", "Meta",
                          "Marcos concluídos", "Atividades em aberto"], linhas)

    def _comps():
        return Competency.query.order_by(Competency.year.desc(),
                                         Competency.month.desc()).all()

    def _projects():
        return Project.query.order_by(Project.name).all()


# --------------------------------------------------------------------------
# Coercoes seguras
# --------------------------------------------------------------------------
def _int(v):
    try:
        return int(v) if v not in (None, "", "None") else None
    except (ValueError, TypeError):
        return None


def _date(v):
    if not v:
        return None
    try:
        return datetime.strptime(v, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return None


# --------------------------------------------------------------------------
# Gantt dos projetos (calculado no servidor; renderizado em CSS puro)
# --------------------------------------------------------------------------
MESES_ABREV = ["", "jan", "fev", "mar", "abr", "mai", "jun",
               "jul", "ago", "set", "out", "nov", "dez"]


def _month_start(d):
    return d.replace(day=1)


def _month_end(d):
    if d.month == 12:
        return d.replace(day=31)
    return d.replace(month=d.month + 1, day=1) - timedelta(days=1)


def build_gantt(projects, prog=None, today=None):
    """Monta a linha do tempo dos projetos.

    Retorna None quando nenhum projeto tem data (nao ha o que desenhar).
    Cada barra traz posicao/largura em % e os marcos posicionados na escala.
    """
    today = today or fuso.hoje()
    prog = prog or {}
    dated = [p for p in projects if p.start_date or p.target_date]
    if not dated:
        return None

    # janela: da menor data ao maior prazo (inclui os marcos e o dia de hoje)
    pontos = []
    for p in dated:
        for d in (p.start_date, p.target_date):
            if d:
                pontos.append(d)
        for m in p.milestones:
            if m.due_date:
                pontos.append(m.due_date)
    pontos.append(today)
    ini = _month_start(min(pontos))
    fim = _month_end(max(pontos))
    total = (fim - ini).days + 1
    if total <= 0:
        return None

    def pct(d):
        return max(0.0, min(100.0, (d - ini).days / total * 100.0))

    # cabecalho por mes
    meses = []
    cur = ini
    while cur <= fim:
        f = _month_end(cur)
        meses.append({
            "label": f"{MESES_ABREV[cur.month]}/{str(cur.year)[2:]}",
            "left": pct(cur),
            "width": ((min(f, fim) - cur).days + 1) / total * 100.0,
            "is_current": cur.year == today.year and cur.month == today.month,
        })
        cur = f + timedelta(days=1)

    linhas = []
    for p in projects:
        ini_p = p.start_date or p.target_date
        fim_p = p.target_date or p.start_date
        if not ini_p:
            linhas.append({"p": p, "sem_data": True, "marcos": []})
            continue
        if fim_p < ini_p:
            ini_p, fim_p = fim_p, ini_p
        left = pct(ini_p)
        width = max(1.2, pct(fim_p) - left)   # largura minima p/ ficar visivel
        info = prog.get(p.id, {})
        marcos = []
        for m in p.milestones:
            if not m.due_date:
                continue
            marcos.append({
                "m": m, "left": pct(m.due_date),
                "atrasado": (not m.done) and m.due_date < today,
            })
        linhas.append({
            "p": p, "sem_data": False, "left": left, "width": width,
            "pct": info.get("pct", 0), "done": info.get("done", 0),
            "total": info.get("total", 0),
            "atrasado": bool(p.target_date and p.target_date < today
                             and p.status not in ("concluido", "cancelado")),
            "marcos": marcos,
        })

    return {"meses": meses, "linhas": linhas,
            "hoje_left": pct(today) if ini <= today <= fim else None,
            "inicio": ini, "fim": fim}


def build_activity_timeline(acts, color="#1d5da8", today=None):
    """Linha do tempo das AÇÕES de um projeto — cada barra é uma atividade.

    Mesma estrutura do build_gantt (meses + linhas) para reusar o CSS do Gantt.
    A barra vai do início (ou criação) até o prazo; cor pela situação.
    """
    today = today or fuso.hoje()
    dated = [a for a in acts if a.due_date or a.start_date]
    if not dated:
        return None
    pontos = []
    for a in dated:
        for d in (a.start_date, a.due_date):
            if d:
                pontos.append(d)
    pontos.append(today)
    ini = _month_start(min(pontos))
    fim = _month_end(max(pontos))
    total = (fim - ini).days + 1
    if total <= 0:
        return None

    def pct(d):
        return max(0.0, min(100.0, (d - ini).days / total * 100.0))

    meses = []
    cur = ini
    while cur <= fim:
        f = _month_end(cur)
        meses.append({"label": f"{MESES_ABREV[cur.month]}/{str(cur.year)[2:]}",
                      "left": pct(cur),
                      "width": ((min(f, fim) - cur).days + 1) / total * 100.0,
                      "is_current": cur.year == today.year and cur.month == today.month})
        cur = f + timedelta(days=1)

    linhas = []
    for a in acts:
        ini_a = a.start_date or a.due_date
        fim_a = a.due_date or a.start_date
        if not ini_a:
            linhas.append({"a": a, "sem_data": True})
            continue
        if fim_a < ini_a:
            ini_a, fim_a = fim_a, ini_a
        concluida = a.status == "concluida"
        atrasada = (not concluida) and a.due_date and a.due_date < today
        cor = "#4caf7d" if concluida else ("#d9534f" if atrasada else color)
        left = pct(ini_a)
        linhas.append({"a": a, "sem_data": False, "left": left,
                       "width": max(1.2, pct(fim_a) - left), "cor": cor,
                       "atrasada": atrasada, "concluida": concluida})
    return {"meses": meses, "linhas": linhas,
            "hoje_left": pct(today) if ini <= today <= fim else None,
            "inicio": ini, "fim": fim}
